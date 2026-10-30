import csv
from pathlib import Path
from data_ingestor import Data

class DataExporter():
    def __init__(self):
        self.path : Path = Path() # .csv, .txt or .xlsx file to write to
        self.params : list[tuple[str, str]] = [] # run parameters, written as '#' comment lines (csv / txt) or on a "params" sheet (excel)
        self.vals : list[int] = [] # the N of the val_N written as columns, in this order
        self._file = None
        self._writer = None
        self._workbook = None
        self._sheet = None

    def __enter__(self) -> "DataExporter":
        self.begin_writing()
        return self

    def __exit__(self, *_) -> None:
        self.end_writing()

    def begin_writing(self):
        """
            creates the file and writes the parameters and the header, raises ValueError if the file type is not supported.
        """

        suffix = self.path.suffix.lower()

        if suffix in [".csv", ".txt"]:
            self._file = open(self.path, "w", newline="", encoding="utf-8")

            # parameters as '#' comment lines (pandas: read_csv(..., comment="#"))
            for key, value in self.params:
                self._file.write(f"# {key}: {value}\n")
            if self.params:
                self._file.write("\n")

            if suffix == ".csv":
                self._writer = csv.writer(self._file)
                self._writer.writerow(self._header())
            return

        if suffix == ".xlsx":
            import openpyxl
            self._workbook = openpyxl.Workbook(write_only=True)
            self._sheet = self._workbook.create_sheet("data")
            self._sheet.append(self._header())
            if self.params:
                params_sheet = self._workbook.create_sheet("params")
                for key, value in self.params:
                    params_sheet.append([key, value])
            return

        raise ValueError(f"the specified output is not supported, valid types: .csv, .txt, .xlsx: {self.path}")

    def write(self, data : Data):
        if self._writer is not None:
            self._write_csv(data)
        elif self._file is not None:
            self._write_txt(data)
        elif self._sheet is not None:
            self._write_excel(data)
        else:
            raise RuntimeError("begin_writing must be called before write")

    def end_writing(self):
        """
            closes the file, excel files are saved here (write only workbooks can not be saved incrementally).
        """

        if self._file is not None:
            self._file.close()
        if self._workbook is not None:
            self._workbook.save(self.path)

        self._file = None
        self._writer = None
        self._workbook = None
        self._sheet = None

    def _header(self) -> list[str]:
        return ["file", *(f"val_{n}" for n in self.vals), "vocals_processed", "vocals_start", "vocals_end", "processing_time", "finished_at"]

    def _write_csv(self, data : Data):
        v_s, v_e = (f"{data.v_s:.3f}", f"{data.v_e:.3f}") if data.v_p else ("", "")
        self._writer.writerow([str(data.file_path), *(getattr(data, f"val_{n}") for n in self.vals), str(data.v_p), v_s, v_e, f"{data.processing_time:.2f}", data.finished_at])
        self._file.flush()

    def _write_excel(self, data : Data):
        v_s, v_e = (round(data.v_s, 3), round(data.v_e, 3)) if data.v_p else (None, None)
        self._sheet.append([str(data.file_path), *(getattr(data, f"val_{n}") for n in self.vals), data.v_p, v_s, v_e, round(data.processing_time, 2), data.finished_at])

    def _write_txt(self, data : Data):
        lines = [f"file: {data.file_path}", *(f"val_{n}: {getattr(data, f'val_{n}')}" for n in self.vals), f"vocals_processed: {data.v_p}"]
        if data.v_p:
            lines += [f"vocals_start: {data.v_s:.3f}", f"vocals_end: {data.v_e:.3f}"]
        lines += [f"processing_time: {data.processing_time:.2f}", f"finished_at: {data.finished_at}"]
        self._file.write("\n".join(lines) + "\n\n")
        self._file.flush()

if __name__ == "__main__":
    import sys

    d1 = Data()
    d1.file_path = Path("a.mp3")
    d1.val_1 = "10"
    d1.v_p = True
    d1.v_s = 12.3456
    d1.v_e = 200.5
    d2 = Data()
    d2.file_path = Path("b.mp3")

    for output in sys.argv[1:] or ["output.csv"]:
        exporter = DataExporter()
        exporter.path = Path(output)
        exporter.params = [("version", "test")]
        exporter.vals = [1]
        with exporter:
            exporter.write(d1)
            exporter.write(d2)
